import tkinter as tk
from tkinter import filedialog, messagebox
import customtkinter as ctk
import os
import json
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

CONFIG_FILE = "opciones_config_mtto.json"

def cargar_opciones():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r") as f:
                return json.load(f)
        except:
            return {"maquinas": [], "usuarios": []}
    return {"maquinas": [], "usuarios": []}

def guardar_opciones(data):
    with open(CONFIG_FILE, "w") as f:
        json.dump(data, f)

class App(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.opciones = cargar_opciones()
        ctk.set_appearance_mode("Dark")

        self.title("REPORTE DE MTTO - Registro Múltiple")
        self.geometry("420x520")
        self.grid_columnconfigure(0, weight=1)

        self.crear_seccion("Máquina", 0, "maquinas")
        self.label_detalle = ctk.CTkLabel(self, text="Detalle del mantenimiento", font=("Arial", 12, "bold"))
        self.label_detalle.grid(row=2, column=0, padx=20, pady=(10, 2), sticky="w")
        self.textbox_detalle = ctk.CTkTextbox(self, height=100)
        self.textbox_detalle.grid(row=3, column=0, padx=20, pady=5, sticky="nsew")
        self.crear_seccion("Usuario", 4, "usuarios")
        
        self.label_foto = ctk.CTkLabel(self, text="Evidencia Fotográfica (Varias)", font=("Arial", 14, "bold"))
        self.label_foto.grid(row=6, column=0, padx=20, pady=(10, 2), sticky="w")
        
        self.btn_foto = ctk.CTkButton(self, text="Examinar imágenes...", fg_color="gray", command=self.cargar_fotos)
        self.btn_foto.grid(row=7, column=0, padx=20, pady=5, sticky="ew")
        
        self.label_contador_fotos = ctk.CTkLabel(self, text="0 fotos seleccionadas", font=("Arial", 10), text_color="gray")
        self.label_contador_fotos.grid(row=8, column=0, padx=20, pady=(0, 5), sticky="w")
        
        self.rutas_fotos = []

        self.save_button = ctk.CTkButton(self, text="Guardar en Word (REPORTE DE MTTO)", fg_color="green", command=self.save_to_word)
        self.save_button.grid(row=9, column=0, padx=20, pady=(20, 20), sticky="ew")

    def crear_seccion(self, titulo, row_idx, key):
        lbl = ctk.CTkLabel(self, text=titulo, font=("Arial", 14, "bold"))
        lbl.grid(row=row_idx, column=0, padx=20, pady=(15, 2), sticky="w")
        frame = ctk.CTkFrame(self, fg_color="transparent")
        frame.grid(row=row_idx+1, column=0, padx=20, pady=5, sticky="ew")
        frame.grid_columnconfigure(0, weight=1)
        entry = ctk.CTkEntry(frame, placeholder_text=f"Seleccione {titulo.lower()}")
        entry.grid(row=0, column=0, sticky="ew", padx=(0, 5))
        if key == "maquinas": self.entry_maquina = entry
        else: self.entry_usuario = entry
        btn_menu = ctk.CTkButton(frame, text="▼", width=40, fg_color="#2b2b2b", command=lambda: self.mostrar_menu(entry, key, btn_menu))
        btn_menu.grid(row=0, column=1, padx=(0, 5))
        btn_add = ctk.CTkButton(frame, text="Add", width=50, fg_color="#1f538d", command=lambda: self.agregar_item(entry, key))
        btn_add.grid(row=0, column=2, padx=(0, 3))
        btn_del = ctk.CTkButton(frame, text="Del", width=50, fg_color="#a83232", command=lambda: self.eliminar_item(entry, key))
        btn_del.grid(row=0, column=3)

    def mostrar_menu(self, entry, key, btn_widget):
        menu = tk.Menu(self, tearoff=0, bg="#2b2b2b", fg="white", font=("Arial", 14))
        for item in self.opciones[key]:
            menu.add_command(label=item, command=lambda val=item: self.seleccionar(entry, val))
        menu.tk_popup(btn_widget.winfo_rootx(), btn_widget.winfo_rooty() + btn_widget.winfo_height())

    def seleccionar(self, entry, valor): 
        entry.delete(0, tk.END)
        entry.insert(0, valor)

    def agregar_item(self, entry, key):
        nueva = entry.get().strip()
        if nueva and nueva not in self.opciones[key]: 
            self.opciones[key].append(nueva)
            guardar_opciones(self.opciones)
            entry.delete(0, tk.END)

    def eliminar_item(self, entry, key):
        valor = entry.get().strip()
        if valor in self.opciones[key]: 
            self.opciones[key].remove(valor)
            guardar_opciones(self.opciones)
            entry.delete(0, tk.END)

    def cargar_fotos(self):
        archivos = filedialog.askopenfilename(
            multiple=True,
            filetypes=[("Imágenes", "*.jpg *.png *.jpeg")]
        )
        if archivos: 
            self.rutas_fotos = list(archivos)
            num_fotos = len(self.rutas_fotos)
            self.btn_foto.configure(text=f"{num_fotos} Fotos Seleccionadas", fg_color="green")
            self.label_contador_fotos.configure(text=f"{num_fotos} archivos listos", text_color="#4CAF50")
        else:
            self.rutas_fotos = []
            self.btn_foto.configure(text="Examinar imágenes...", fg_color="gray")
            self.label_contador_fotos.configure(text="0 fotos seleccionadas", text_color="gray")

    def save_to_word(self):
        maquina = self.entry_maquina.get().strip()
        usuario = self.entry_usuario.get().strip()
        detalle = self.textbox_detalle.get("1.0", tk.END).strip()
        
        if not maquina or not usuario or not detalle:
            messagebox.showwarning("Campos incompletos", "Por favor completa la máquina, el usuario y el detalle antes de guardar.")
            return

        # Limpiar caracteres prohibidos para nombres de archivos en Windows
        maquina_limpia = "".join(c for c in maquina if c.isalnum() or c in (' ', '_', '-')).strip()
        usuario_limpio = "".join(c for c in usuario if c.isalnum() or c in (' ', '_', '-')).strip()
        
        ahora = datetime.now()
        fecha_str = ahora.strftime("%d/%m/%Y")
        hora_str = ahora.strftime("%H:%M")
        fecha_archivo = ahora.strftime("%d-%m-%Y_%H-%M")

        filename = f"{maquina_limpia} - {usuario_limpio} - {fecha_archivo}.docx"
        
        doc = Document()
        p_title = doc.add_paragraph()
        run_title = p_title.add_run("REPORTE GENERAL DE MANTENIMIENTO")
        run_title.font.name = 'Arial'
        run_title.font.size = Pt(16)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(31, 83, 141)
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()

        p_sub = doc.add_paragraph()
        run_sub = p_sub.add_run(f"Mantenimiento en Máquina: {maquina} ({fecha_str} - {hora_str})")
        run_sub.font.name = 'Arial'
        run_sub.font.size = Pt(14)
        run_sub.font.bold = True
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Tabla de Datos Principales
        table = doc.add_table(rows=4, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        datos_principales = [
            ("Fecha / Hora", f"{fecha_str} - {hora_str}"),
            ("Máquina", maquina),
            ("Técnico / Usuario", usuario),
            ("Detalle de MTTO", detalle)
        ]

        for i, (campo, valor) in enumerate(datos_principales):
            row = table.rows[i]
            cell_0 = row.cells[0]
            cell_0.text = campo
            cell_0.paragraphs[0].runs[0].font.bold = True
            cell_1 = row.cells[1]
            cell_1.text = valor

        doc.add_paragraph()

        # Lógica de 2 imágenes por página con espacio individual
        if self.rutas_fotos:
            num_total_fotos = len(self.rutas_fotos)
            
            for i in range(0, num_total_fotos, 2):
                doc.add_page_break() 
                
                p_img_title = doc.add_paragraph()
                run_img_title = p_img_title.add_run("Evidencia Fotográfica del Mantenimiento:")
                run_img_title.font.name = 'Arial'
                run_img_title.font.size = Pt(12)
                run_img_title.font.bold = True
                
                doc.add_paragraph() 

                # Primera imagen
                ruta_1 = self.rutas_fotos[i]
                if os.path.exists(ruta_1):
                    p_num1 = doc.add_paragraph()
                    run_num1 = p_num1.add_run(f"Evidencia {i+1} de {num_total_fotos}")
                    run_num1.font.size = Pt(9)
                    run_num1.font.color.rgb = RGBColor(100, 100, 100)
                    p_num1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    
                    p_img1 = doc.add_paragraph()
                    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img1.add_run().add_picture(ruta_1, width=Inches(4.5))

                # Espaciador
                espaciador = doc.add_paragraph()
                espaciador.paragraph_format.space_after = Pt(18)

                # Segunda imagen (si existe)
                if i + 1 < num_total_fotos:
                    ruta_2 = self.rutas_fotos[i+1]
                    if os.path.exists(ruta_2):
                        p_num2 = doc.add_paragraph()
                        run_num2 = p_num2.add_run(f"Evidencia {i+2} de {num_total_fotos}")
                        run_num2.font.size = Pt(9)
                        run_num2.font.color.rgb = RGBColor(100, 100, 100)
                        p_num2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        
                        p_img2 = doc.add_paragraph()
                        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img2.add_run().add_picture(ruta_2, width=Inches(4.5))

        try:
            doc.save(filename)
            messagebox.showinfo("Éxito", f"Se guardó correctamente\n\nArchivo:\n{filename}")
        except PermissionError:
            messagebox.showerror("Error de permisos", f"No se pudo guardar el archivo '{filename}'.\n\nAsegúrate de que no esté abierto en Microsoft Word.")
            return
        
        # Limpiar formulario
        self.entry_maquina.delete(0, tk.END)
        self.entry_usuario.delete(0, tk.END)
        self.textbox_detalle.delete("1.0", tk.END)
        self.btn_foto.configure(text="Examinar imágenes...", fg_color="gray")
        self.label_contador_fotos.configure(text="0 fotos seleccionadas", text_color="gray")
        self.rutas_fotos = []

if __name__ == "__main__":
    app = App()
    app.mainloop()